import re, json
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL

TOL = 0.05
DATE = r"^\d{1,2} [A-Z][a-z]+\.? \d{4}$"
CITE = r"\([^()]*\d[^()]*\)"
QUOTE = r'“([^“”]+)”|"([^"]+)"'

# (key, label, needs .docx)
CHECKS = [
    ("margins", '1" margins on all sides', True),
    ("font", "Times New Roman, 12 pt", True),
    ("header", "Last name + page number header", True),
    ("heading", "Heading block with date as Day Month Year", False),
    ("title", "Title centered, no bold/underline", True),
    ("spacing", "Double spaced", True),
    ("indent", '0.5" paragraph indents', True),
    ("cites", "In-text citations like (Author 42)", False),
    ("q_marks", "Quotation marks balanced and consistent", False),
    ("q_punct", "Punctuation around quotes (period after citation)", False),
    ("q_cite", "Quotes followed by a citation", False),
    ("q_block", "Long quotes set as block quotes", False),
    ("wc", "Works Cited page and heading", False),
    ("wc_page", "Works Cited on its own page", True),
    ("wc_indent", 'Works Cited hanging indent (0.5")', True),
    ("wc_alpha", "Works Cited in alphabetical order", False),
    ("wc_url", "URLs without https://", False),
]

def ps(p, attr):
    v = getattr(p.paragraph_format, attr); s = p.style
    while v is None and s is not None:
        v = getattr(s.paragraph_format, attr); s = s.base_style
    return v

def rf(p, r, attr):
    v = getattr(r.font, attr); s = p.style
    while v is None and s is not None:
        v = getattr(s.font, attr); s = s.base_style
    return v

def inch(x): return None if x is None else x / 914400
def fancy(p): return any(r.bold or r.underline for r in p.runs) or bool(p.style.font.bold)
def new_page(P, i):
    x = P[i]._p.xml
    return bool(ps(P[i], "page_break_before")) or 'w:type="page"' in x or "lastRenderedPageBreak" in x or (i > 0 and 'w:type="page"' in P[i-1]._p.xml)

def analyze(texts, doc=None):
    P = doc.paragraphs if doc else None
    out, seen = [], {"quotes": 0, "urls": 0, "wc": False}
    def add(key, level, msg, para=None, frag=None):
        out.append(dict(key=key, level=level, msg=msg, para=para, frag=frag))
    ne = [i for i, t in enumerate(texts) if t.strip()]
    if len(ne) < 6:
        add("heading", "ERROR", "Too short to check. Is this the full essay?")
        return out, []

    if doc:
        s = doc.sections[0]
        for n, v in [("Top", s.top_margin), ("Bottom", s.bottom_margin), ("Left", s.left_margin), ("Right", s.right_margin)]:
            if v is None or abs(inch(v) - 1) > TOL:
                add("margins", "ERROR", f'{n} margin is {"unknown" if v is None else f"{inch(v):.2f}"}". MLA needs 1" on all sides.')
        fonts, sizes = set(), set()
        for p in P:
            for r in p.runs:
                if r.text.strip():
                    fonts.add(rf(p, r, "name")); z = rf(p, r, "size"); sizes.add(z.pt if z else None)
        bad = sorted(f for f in fonts if f and f != "Times New Roman")
        if bad: add("font", "ERROR", f"Font is {', '.join(bad)}. Use Times New Roman.")
        bs = sorted(z for z in sizes if z and z != 12)
        if bs: add("font", "ERROR", f"Font size is {', '.join(str(z) for z in bs)} pt. Use 12 pt.")
        hx = "".join(p._p.xml for p in s.header.paragraphs)
        ht = " ".join(p.text for p in s.header.paragraphs).strip()
        if not ht and "PAGE" not in hx:
            add("header", "ERROR", "No header. Add your last name + page number, top right (e.g. Smith 1).")
        else:
            if "PAGE" not in hx: add("header", "ERROR", "Header has no automatic page number (Insert > Page Number).")
            if not re.search(r"[A-Za-z]{2,}", ht): add("header", "ERROR", "Header is missing your last name before the page number.")
            if s.header.paragraphs[0].alignment != AL.RIGHT: add("header", "WARN", "Header should be right-aligned.")

    h = ne[:4]
    if not re.match(DATE, texts[h[3]].strip()):
        add("heading", "ERROR", "4th line should be the date as Day Month Year (e.g. 3 October 2026).", h[3])
    if doc:
        for i in h:
            if P[i].alignment in (AL.CENTER, AL.RIGHT):
                add("heading", "ERROR", "Heading lines (name/instructor/course/date) should be left-aligned.", i)
    ti = ne[4]
    if doc:
        if P[ti].alignment != AL.CENTER: add("title", "ERROR", "Title should be centered.", ti)
        if fancy(P[ti]): add("title", "ERROR", "Title shouldn't be bold or underlined.", ti)
    if texts[ti].strip().endswith("."): add("title", "WARN", "Title shouldn't end with a period.", ti)

    wc = next((i for i in ne if texts[i].strip().lower() in ("works cited", "work cited")), None)
    end = wc if wc is not None else len(texts)
    cited = False
    body = [i for i in ne if ti < i < end]
    allt = " ".join(texts[i] for i in body)
    if '"' in allt and ("“" in allt or "”" in allt):
        add("q_marks", "WARN", 'Mixes straight (") and curly (“ ”) quotation marks. Pick one style.')

    for i in body:
        t = texts[i]
        block = False
        if doc:
            li, fi = ps(P[i], "left_indent"), ps(P[i], "first_line_indent")
            block = li is not None and abs(inch(li) - .5) <= TOL and (fi is None or abs(inch(fi)) <= TOL)
        if block:
            seen["quotes"] += 1
            if re.match(r'\s*["“]', t): add("q_block", "ERROR", "Block quotes don't use quotation marks.", i, t.strip()[:25])
            if re.search(r"\)[.!?]\s*$", t): add("q_punct", "ERROR", "In a block quote the period goes before the citation, not after.", i, t.strip()[-25:])
        if doc and len(t) > 120:
            ls = ps(P[i], "line_spacing")
            if not isinstance(ls, float) or abs(ls - 2) > .05: add("spacing", "ERROR", "Not double spaced.", i)
            if not block:
                fi = ps(P[i], "first_line_indent")
                if fi is None or abs(inch(fi) - .5) > TOL: add("indent", "ERROR", 'Missing 0.5" first-line indent.', i)
            sa = ps(P[i], "space_after")
            if sa and sa > Pt(1): add("spacing", "ERROR", "Extra space after paragraph. Set it to 0 pt.", i)
        for c in re.findall(r"\(([^()]*\d[^()]*)\)", t):
            cited = True; f = f"({c})"
            if re.search(r"\b(p|pp|pg|page)\b\.?", c, re.I): add("cites", "ERROR", 'Drop "p."/"pg". Just (Author 42).', i, f)
            elif re.match(r"[A-Z][\w'\- ]*,?\s*(19|20)\d{2}", c): add("cites", "ERROR", "Looks like APA (has a year). MLA uses (Author page).", i, f)
            elif re.search(r"[A-Za-z],\s*\d", c): add("cites", "ERROR", "No comma between author and page number.", i, f)

        # quotation marks
        if t.count('"') % 2 or t.count("“") != t.count("”"):
            add("q_marks", "WARN", "Odd number of quotation marks. Is one missing?", i); continue
        for m in re.finditer(QUOTE, t):
            q = m.group(1) or m.group(2); words = len(q.split()); seen["quotes"] += 1
            after = t[m.end():]
            cm = re.match(r"\s*" + CITE, after)
            if cm:
                frag = t[max(m.start(), m.end() - 25): m.end() + cm.end()]
                if q.rstrip()[-1:] in ".,":
                    add("q_punct", "ERROR", "Period/comma goes after the citation, not inside the quote: \"…quote\" (Smith 42).", i, frag)
                rest = after[cm.end():]
                if rest[:1] not in (".", ",", ";", ":", "?", "!", ")") and (not rest.strip() or re.match(r"\s+[A-Z]", rest)):
                    add("q_punct", "WARN", "End the sentence with a period after the citation.", i, frag)
            elif words >= 8 and not re.search(CITE, re.split(r"[.!?]\s+(?=[A-Z])", after, 1)[0]):
                add("q_cite", "WARN", "This quote has no citation after it.", i, t[max(m.start(), m.end() - 25): m.end()])
            if words > 50:
                add("q_block", "WARN", f"{words}-word quote. Over about 4 lines, use a block quote (indent 0.5\", no quotation marks).", i, t[m.start(): m.start() + 40])
    if not cited: add("cites", "WARN", "No in-text citations found. MLA expects (Author page), e.g. (Smith 42).")

    if wc is None:
        add("wc", "ERROR", 'No "Works Cited" page found.')
    else:
        seen["wc"] = True
        if texts[wc].strip() != "Works Cited": add("wc", "ERROR", 'Heading should read exactly "Works Cited".', wc)
        if doc:
            if P[wc].alignment != AL.CENTER: add("wc", "ERROR", "Works Cited heading should be centered.", wc)
            if fancy(P[wc]): add("wc", "ERROR", "Works Cited heading shouldn't be bold or underlined.", wc)
            if not new_page(P, wc): add("wc_page", "ERROR", "Works Cited should start on its own page.", wc)
        ents = [i for i in ne if i > wc]
        if not ents: add("wc", "ERROR", "Works Cited has no entries.", wc)
        prev = ""
        for i in ents:
            t = texts[i]
            if doc:
                li, fi = ps(P[i], "left_indent"), ps(P[i], "first_line_indent")
                if not (fi is not None and fi < 0 and li is not None and abs(inch(li) - .5) <= TOL):
                    add("wc_indent", "ERROR", 'Needs a 0.5" hanging indent.', i)
            k = re.sub(r"^(the|a|an)\s+", "", t.strip().lower())
            if k < prev: add("wc_alpha", "ERROR", "Out of alphabetical order.", i)
            prev = k
            m = re.search(r"https?://\S+", t)
            if re.search(r"(https?://|www\.)\S+", t): seen["urls"] += 1
            if m: add("wc_url", "WARN", "MLA 9 drops https://. Start at www. or the domain.", i, m.group(0))

    checklist = []
    for key, label, needs_doc in CHECKS:
        hits = [x for x in out if x["key"] == key]
        if hits:
            st = "fail" if any(x["level"] == "ERROR" for x in hits) else "warn"
        elif needs_doc and not doc: st = "skip"
        elif key.startswith("q_") and key != "q_marks" and not seen["quotes"]: st = "na"
        elif key.startswith("wc_") and not seen["wc"]: st = "na"
        elif key == "wc_url" and not seen["urls"]: st = "na"
        else: st = "pass"
        checklist.append(dict(label=label, status=st, count=len(hits)))
    return out, checklist

def pack(texts, doc=None):
    issues, checklist = analyze(texts, doc)
    return json.dumps(dict(paras=texts, issues=issues, checklist=checklist, formatting=bool(doc)))

def run_docx(path):
    doc = Document(path)
    return pack([p.text for p in doc.paragraphs], doc)

def run_text(s):
    return pack([l for l in s.splitlines() if l.strip()])
